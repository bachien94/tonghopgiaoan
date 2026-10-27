import streamlit as st
import datetime
import io
import copy
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Cấu hình trang rộng để hiển thị bảng TKB thoáng hơn
st.set_page_config(page_title="Tổng Hợp Giáo Án Theo Tuần", page_icon="📚", layout="wide")

st.markdown("## 📊 Phần Mềm Tổng Hợp Giáo Án Theo Tuần")

# --- PHẦN HEADER: TUẦN VÀ NGÀY ---
col_h1, col_h2 = st.columns([1, 3])
with col_h1:
    tuan_so = st.number_input("Tuần số:", min_value=1, max_value=52, value=1)
with col_h2:
    ngay_dau_tuan = st.date_input("Ngày Thứ Hai đầu tuần (dd/mm/yyyy):", value=datetime.date.today())

st.markdown("---")
st.markdown("### Bảng Thời Khóa Biểu *(Sáng: Xanh dương | Chiều: Vàng kem)*")

# Danh sách môn học mẫu
danh_sach_mon = [
    "-- Chọn môn học --", "Tin học", "Toán", "Tiếng Việt", "Đạo đức", 
    "Tự nhiên và Xã hội", "Khoa học", "Lịch sử và Địa lý", 
    "Mỹ thuật", "Âm nhạc", "Thể dục", "Hoạt động trải nghiệm", "Tiếng Anh"
]

danh_sach_lop = ["-- Chọn lớp --", "Lớp 1", "Lớp 2", "Lớp 3", "Lớp 4", "Lớp 5"]

# Các ngày trong tuần
cac_thu = ["Thứ Hai", "Thứ Ba", "Thứ Tư", "Thứ Năm", "Thứ Sáu"]

# Mapping số ngày cộng thêm từ Thứ Hai đầu tuần
ngay_offset = {
    "Thứ Hai": 0,
    "Thứ Ba": 1,
    "Thứ Tư": 2,
    "Thứ Năm": 3,
    "Thứ Sáu": 4
}

# Các tiết học (Sáng từ tiết 1-5, Chiều từ tiết 1-3)
dieu_phoi_tiet = [
    ("Buổi Sáng", "Tiết 1"),
    ("Buổi Sáng", "Tiết 2"),
    ("Buổi Sáng", "Tiết 3"),
    ("Buổi Sáng", "Tiết 4"),
    ("Buổi Sáng", "Tiết 5"),
    ("Buổi Chiều", "Tiết 1"),
    ("Buổi Chiều", "Tiết 2"),
    ("Buổi Chiều", "Tiết 3"),
]

# Vẽ tiêu đề bảng (Header)
cols_header = st.columns([1.5] + [2] * len(cac_thu))
with cols_header[0]:
    st.markdown("**Buổi / Tiết**")
for i, thu in enumerate(cac_thu):
    with cols_header[i + 1]:
        st.markdown(f"**{thu}**")

st.markdown("<hr style='margin: 5px 0;'>", unsafe_allow_html=True)

# Lưu trữ dữ liệu nhập liệu của các ô
tkb_data = {}

# Vẽ các dòng tiết học
for buoi, tiet in dieu_phoi_tiet:
    row_cols = st.columns([1.5] + [2] * len(cac_thu))
    
    # Cột tên buổi/tiết (Phân màu Sáng / Chiều)
    with row_cols[0]:
        if "Sáng" in buoi:
            st.markdown(f"<div style='background-color: #e6f2ff; padding: 10px; border-radius: 5px; font-weight: bold; text-align: center;'>{buoi}<br>{tiet}</div>", unsafe_allow_html=True)
        else:
            st.markdown(f"<div style='background-color: #fff9e6; padding: 10px; border-radius: 5px; font-weight: bold; text-align: center;'>{buoi}<br>{tiet}</div>", unsafe_allow_html=True)
    
    # Các ô nhập liệu cho từng thứ trong tuần
    for i, thu in enumerate(cac_thu):
        with row_cols[i + 1]:
            key_prefix = f"{thu}_{buoi}_{tiet}"
            
            # Chọn môn học
            mon = st.selectbox(f"Môn {key_prefix}", danh_sach_mon, label_visibility="collapsed", key=f"mon_{key_prefix}")
            # Chọn lớp
            lop = st.selectbox(f"Lớp {key_prefix}", danh_sach_lop, label_visibility="collapsed", key=f"lop_{key_prefix}")
            # Chọn file giáo án
            file_ga = st.file_uploader(f"File {key_prefix}", type=["docx", "pdf"], label_visibility="collapsed", key=f"file_{key_prefix}")
            
            # Lưu lại thông tin nếu người dùng đã chọn môn
            if mon != "-- Chọn môn học --":
                tkb_data[key_prefix] = {
                    "thu": thu, 
                    "buoi": buoi, 
                    "tiet": tiet, 
                    "mon": mon, 
                    "lop": lop, 
                    "file": file_ga
                }

    st.markdown("<hr style='margin: 5px 0; border-color: #f0f2f6;'>", unsafe_allow_html=True)

# --- HÀM HỖ TRỢ GHÉP FILE WORD ---
def generate_master_docx(tkb_dict, tuan, ngay_bd):
    master_doc = Document()
    
    # Cấu hình lề trang (Margins) theo yêu cầu: Trái 3cm, Phải 1.5cm, Trên 2cm, Dưới 2cm
    sections = master_doc.sections
    for section in sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)
    
    # Tiêu đề chung của file tổng hợp (Trang bìa / Mở đầu)
    p_title = master_doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(f"TỔNG HỢP GIÁO ÁN - TUẦN {tuan}")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(16)
    r_title.bold = True
    
    master_doc.add_page_break() # Ngắt trang sau tiêu đề tuần để sang ngày mới
    
    last_written_day = None

    # Duyệt theo thứ tự ngày trong tuần
    for thu in cac_thu:
        ngay_cu_the = ngay_bd + datetime.timedelta(days=ngay_offset[thu])
        tiet_trong_ngay = [(b, t) for b, t in dieu_phoi_tiet if f"{thu}_{b}_{t}" in tkb_dict]
        
        if tiet_trong_ngay:
            # Nếu sang ngày mới (khác ngày đầu tiên), tự động sang trang mới hoàn toàn
            if last_written_day is not None:
                master_doc.add_page_break()
                
            # Ghi thông tin Thứ, ngày, tháng cho ngày mới này
            p_date = master_doc.add_paragraph()
            p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_date = p_date.add_run(f"{thu}, ngày {ngay_cu_the.day} tháng {ngay_cu_the.month} năm {ngay_cu_the.year}")
            r_date.font.name = 'Times New Roman'
            r_date.font.size = Pt(14)
            r_date.italic = True
            
            last_written_day = thu
            
            # Đếm số tiết trong ngày để quản lý việc kẻ vạch ngang giữa các môn
            total_classes_today = len(tiet_trong_ngay)
            current_class_today_idx = 0

            for buoi, tiet in tiet_trong_ngay:
                key = f"{thu}_{buoi}_{tiet}"
                item = tkb_dict[key]
                current_class_today_idx += 1
                
                # 1. TÊN LỚP (Căn giữa, in đậm)
                p_lop = master_doc.add_paragraph()
                p_lop.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_lop = p_lop.add_run(f"{item['lop']}")
                r_lop.font.name = 'Times New Roman'
                r_lop.font.size = Pt(14)
                r_lop.bold = True
                
                # 2. TÊN MÔN HỌC (Viết hoa, căn giữa, in đậm - Không hiển thị số tiết)
                p_mon = master_doc.add_paragraph()
                p_mon.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_mon = p_mon.add_run(f"MÔN: {item['mon'].upper()}")
                r_mon.font.name = 'Times New Roman'
                r_mon.font.size = Pt(14)
                r_mon.bold = True
                
                # 3. NỘI DUNG BÀI HỌC (Chèn đúng nội dung file con vào)
                uploaded_file = item["file"]
                if uploaded_file is not None and uploaded_file.name.endswith(".docx"):
                    try:
                        sub_doc = Document(uploaded_file)
                        anchor_p = master_doc.add_paragraph()
                        
                        for element in sub_doc.element.body:
                            if element.tag.endswith('p') or element.tag.endswith('tbl'):
                                anchor_p._p.addprevious(copy.deepcopy(element))
                                
                        anchor_p._p.getparent().remove(anchor_p._p)
                    except Exception as e:
                        p_err = master_doc.add_paragraph(f"[Lỗi đọc file: {str(e)}]")
                        p_err.runs[0].font.name = 'Times New Roman'
                else:
                    p_none = master_doc.add_paragraph("[Chưa đính kèm file Word giáo án cho tiết này]")
                    p_none.runs[0].font.name = 'Times New Roman'
                    p_none.runs[0].italic = True
                
                # 4. KẺ ĐƯỜNG NGANG KHOẢNG NỬA TRANG SAU MỖI MÔN (Trừ môn cuối cùng trong ngày)
                if current_class_today_idx < total_classes_today:
                    p_line = master_doc.add_paragraph()
                    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r_line = p_line.add_run("--------------------------------------------------")
                    r_line.font.name = 'Times New Roman'
                    r_line.font.size = Pt(12)
                    r_line.bold = True
                
    # Lưu vào bộ nhớ đệm (BytesIO)
    file_stream = io.BytesIO()
    master_doc.save(file_stream)
    file_stream.seek(0)
    return file_stream

# --- PHẦN NÚT XUẤT FILE Ở DƯỚI CÙNG ---
st.markdown("<br>", unsafe_allow_html=True)
col_btn1, col_btn2, col_btn3 = st.columns([1, 2, 1])

with col_btn2:
    if st.button("📥 Xử Lý & Tạo File Word Tổng Hợp", type="primary", use_container_width=True):
        if not tkb_data:
            st.warning("Vui lòng chọn ít nhất một môn học trong thời khóa biểu!")
        else:
            with st.spinner("Đang tổng hợp giáo án, vui lòng chờ trong giây lát..."):
                docx_stream = generate_master_docx(tkb_data, tuan_so, ngay_dau_tuan)
                st.session_state['generated_docx'] = docx_stream.getvalue()
                st.success(f"Đã tổng hợp thành công giáo án cho Tuần {tuan_so}!")

# Hiển thị nút Download trực tiếp khi đã xử lý xong
if 'generated_docx' in st.session_state:
    st.download_button(
        label="📥 Bấm vào đây để tải file Word Tổng Hợp về máy",
        data=st.session_state['generated_docx'],
        file_name=f"Tong_Hop_Giao_An_Tuan_{tuan_so}.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        type="primary",
        use_container_width=True
    )
